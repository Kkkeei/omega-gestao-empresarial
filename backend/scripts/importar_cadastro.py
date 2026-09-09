from pathlib import Path
import re
import sys
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
BACKEND = ROOT / 'backend'
sys.path.insert(0, str(BACKEND))
from app.db.database import conectar_banco, criar_tabelas


def cnpj_limpo(value):
    return re.sub(r"\D", "", str(value or ""))


def main():
    arquivo = ROOT / "dados" / "CADASTRO.xlsx"
    if not arquivo.exists():
        raise SystemExit(f"Arquivo não encontrado: {arquivo}")
    criar_tabelas()
    wb = load_workbook(arquivo, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    inseridas = atualizadas = ignoradas = 0
    conexao = conectar_banco()
    try:
        for nome, cnpj in ws.iter_rows(min_row=2, max_col=2, values_only=True):
            nome = " ".join(str(nome or "").split())
            cnpj = cnpj_limpo(cnpj)
            if not nome or len(cnpj) != 14:
                ignoradas += 1
                continue
            existente = conexao.execute("SELECT id, razao_social FROM empresas WHERE cnpj=?", (cnpj,)).fetchone()
            if existente:
                if existente["razao_social"] != nome:
                    conexao.execute("UPDATE empresas SET razao_social=?, atualizado_em=CURRENT_TIMESTAMP WHERE id=?", (nome, existente["id"]))
                    atualizadas += 1
                continue
            cur = conexao.execute("INSERT INTO empresas (cnpj, razao_social) VALUES (?,?)", (cnpj, nome))
            conexao.execute("INSERT INTO empresa_historicos (empresa_id,tipo_evento,descricao,origem) VALUES (?,?,?,?)", (cur.lastrowid, "IMPORTACAO", "Empresa importada do CADASTRO.xlsx", "IMPORTACAO"))
            inseridas += 1
        conexao.commit()
    finally:
        conexao.close()
        wb.close()
    print(f"Importação concluída: {inseridas} inseridas, {atualizadas} atualizadas, {ignoradas} ignoradas.")


if __name__ == "__main__":
    main()
